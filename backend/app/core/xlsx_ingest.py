from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.core.models import Event

WORKFLOW_LABELS = {
    "AGY_TRADE": "Agency Trade",
    "SL_CREATE_TRADE": "SLT Trade Creation",
    "SL_ALTER_TRADE": "SLT Trade Settlement",
    "CANCEL_TRADE": "SLT Trade Cancellation",
    "SL_CREATE_SUB_ALLOCATION": "SLT Trade Allocation",
    "SL_CREATE_RE_ALLOCATION": "SLT Trade Reallocation",
    "ELEVATION": "SLT Trade Elevation",
}


def _cell(row: dict[str, Any], *names: str) -> Any:
    for name in names:
        if name in row and row[name] not in (None, ""):
            return row[name]
    return None


def _files(value: Any) -> list[str]:
    if value in (None, ""):
        return []
    return [part.strip() for part in str(value).replace("\r", "").split("\n") if part.strip()]


def _dt(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S.%f", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def _num(value: Any) -> float | None:
    if value in (None, ""):
        return None
    return float(value)


def _int(value: Any) -> int:
    if value in (None, ""):
        return 0
    return int(value)


def _headers(sheet) -> list[str]:
    return [str(cell.value).strip() if cell.value is not None else "" for cell in next(sheet.iter_rows(min_row=1, max_row=1))]


def _rows(sheet) -> list[dict[str, Any]]:
    headers = _headers(sheet)
    items: list[dict[str, Any]] = []
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if not any(values):
            continue
        items.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})
    return items


def _event_id(category: str, workflow: str, trade_id: str) -> str:
    return f"trade:{category}:{workflow}:{trade_id}"


def _from_row(row: dict[str, Any], category: str, default_workflow: str) -> Event | None:
    trade_id = str(_cell(row, "Trade ID") or "").strip()
    if not trade_id:
        return None
    workflow = str(_cell(row, "Workflow") or default_workflow)
    request_type = str(_cell(row, "Request Type") or WORKFLOW_LABELS.get(workflow, category + " Trade"))
    return Event(
        id=_event_id(category, workflow, trade_id),
        domain="trade",
        display_id=trade_id,
        group_id=str(_cell(row, "Deal ID") or ""),
        category=category,
        workflow=workflow,
        request_type=request_type,
        status=str(_cell(row, "Status") or "Unmatched"),
        outcome=_cell(row, "Outcome"),
        start=_dt(_cell(row, "Processing Start")),
        end=_dt(_cell(row, "Response Time")),
        duration_seconds=_num(_cell(row, "Duration Seconds")),
        request_count=_int(_cell(row, "Request Count")),
        response_count=_int(_cell(row, "Response Count")),
        request_files=_files(_cell(row, "Request Source Files")),
        response_files=_files(_cell(row, "Response Source Files")),
        request_payload=_cell(row, "Request Payload"),
        response_payload=_cell(row, "Response Payload"),
        tags={},
    )


def ingest_trade_workbook(path: Path) -> list[Event]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    events: list[Event] = []
    for sheet in workbook.worksheets:
        name = (sheet.title or "").strip().lower()
        if name == "agency":
            category, workflow = "Agency", "AGY_TRADE"
        elif name == "slt":
            category, workflow = "SLT", "SL_ALTER_TRADE"
        else:
            continue
        for row in _rows(sheet):
            event = _from_row(row, category, workflow)
            if event:
                events.append(event)
    return events


def find_workbook(folder: Path) -> Path | None:
    direct = folder / "trade_analysis.xlsx"
    if direct.exists():
        return direct
    matches = sorted(folder.rglob("trade_analysis.xlsx"))
    if matches:
        return matches[0]
    parent = folder.parent / "trade_analysis.xlsx"
    if parent.exists():
        return parent
    return None
