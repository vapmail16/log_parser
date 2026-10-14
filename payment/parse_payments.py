from __future__ import annotations

import argparse
import gzip
import json
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    raise SystemExit(1)

HEADER = re.compile(
    r"\[(?P<ts>\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2},\d{3})\s+(?P<thread>[^\]]+)\]"
)
REF = re.compile(
    r'paymentTransactionReference"\s*:\s*"([^"]+)"',
    re.IGNORECASE,
)
TS_FORMAT = "%Y-%m-%d %H:%M:%S,%f"
SEND_HINTS = ("sending", "send ", "sent", "request", "outbound")
RESPONSE_HINTS = ("response", "received", "inbound", "success", "ack", "nack", "reply")
OUTPUT_NAME = "payment_timings.xlsx"
COLUMNS = [
    ("payment_transaction_reference", "Payment Transaction Reference"),
    ("sent_at", "Sent At"),
    ("response_at", "Response At"),
    ("duration_seconds", "Duration Seconds"),
    ("status", "Status"),
    ("source_files", "Source Files"),
]


def load_config(path: Path) -> Path:
    data = json.loads(path.read_text(encoding="utf-8"))
    folder = Path(str(data["log_folder"])).expanduser()
    if not folder.is_absolute():
        folder = (path.parent / folder).resolve()
    return folder.resolve()


def collect_log_files(folder: Path) -> list[Path]:
    files = [
        path
        for path in sorted(folder.rglob("*"))
        if path.is_file() and path.suffix.lower() in {".log", ".gz"}
    ]
    return files


def _open_log(path: Path):
    if path.suffix.lower() == ".gz":
        return gzip.open(path, "rt", encoding="utf-8", errors="ignore")
    return path.open("rt", encoding="utf-8", errors="ignore")


def _parse_ts(raw: str) -> datetime:
    return datetime.strptime(raw, TS_FORMAT)


def _kind(text: str) -> str:
    lower = text.lower()
    if any(hint in lower for hint in RESPONSE_HINTS):
        return "response"
    if any(hint in lower for hint in SEND_HINTS):
        return "send"
    return "other"


def extract_hits(lines: list[str], source: str = "") -> list[dict]:
    hits: list[dict] = []
    current: dict | None = None
    seen: set[str] = set()
    for raw in lines:
        line = raw.rstrip("\n")
        header = HEADER.search(line)
        if header:
            current = {
                "timestamp": _parse_ts(header.group("ts")),
                "text": line,
                "source": source,
            }
            seen = set()
        elif current is not None:
            current["text"] = f"{current['text']}\n{line}"
        blob = current["text"] if current else line
        for ref in REF.findall(blob):
            if current is None or ref in seen:
                continue
            seen.add(ref)
            hits.append(
                {
                    "ref": ref,
                    "timestamp": current["timestamp"],
                    "kind": _kind(blob),
                    "source": current["source"],
                }
            )
    return hits


def summarize(hits: list[dict]) -> list[dict]:
    by_ref: dict[str, list[dict]] = defaultdict(list)
    for hit in hits:
        by_ref[hit["ref"]].append(hit)
    rows: list[dict] = []
    for ref, items in sorted(by_ref.items()):
        items = sorted(items, key=lambda item: item["timestamp"])
        send = next((item for item in items if item["kind"] == "send"), items[0])
        later = [item for item in items if item["timestamp"] >= send["timestamp"] and item is not send]
        response = next((item for item in reversed(later) if item["kind"] == "response"), None)
        if response is None and later:
            response = later[-1] if later[-1]["kind"] != "send" else None
        sent_at = send["timestamp"]
        response_at = response["timestamp"] if response else None
        duration = round((response_at - sent_at).total_seconds(), 3) if response_at else None
        files = sorted({Path(item["source"]).name for item in items if item["source"]})
        rows.append(
            {
                "payment_transaction_reference": ref,
                "sent_at": sent_at.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
                "response_at": response_at.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3] if response_at else None,
                "duration_seconds": duration,
                "status": "Matched" if response_at else "Unmatched",
                "source_files": "\n".join(files),
            }
        )
    return rows


def parse_folder(folder: Path) -> list[dict]:
    hits: list[dict] = []
    for path in collect_log_files(folder):
        with _open_log(path) as handle:
            hits.extend(extract_hits(handle.read().splitlines(), str(path)))
    return summarize(hits)


def write_workbook(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payments"
    headers = [label for _key, label in COLUMNS]
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")
    matched_fill = PatternFill("solid", fgColor="C6EFCE")
    unmatched_fill = PatternFill("solid", fgColor="FCE4D6")
    for row in rows:
        sheet.append([row.get(key) for key, _label in COLUMNS])
    for row_number in range(2, sheet.max_row + 1):
        fill = matched_fill if sheet.cell(row_number, 5).value == "Matched" else unmatched_fill
        for column in range(1, 7):
            sheet.cell(row_number, column).fill = fill
            sheet.cell(row_number, column).alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    widths = [32, 24, 24, 18, 14, 28]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    workbook.save(path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Find payment send and response times in log files.")
    parser.add_argument("log_folder", nargs="?", help="Folder of .log / .log.gz files. Overrides config.json.")
    parser.add_argument(
        "--config",
        default=str(Path(__file__).resolve().parent / "config.json"),
        help="Config file with log_folder.",
    )
    args = parser.parse_args(argv)
    if args.log_folder:
        folder = Path(args.log_folder).expanduser().resolve()
    else:
        folder = load_config(Path(args.config))
    if not folder.is_dir():
        print(f"Log folder does not exist: {folder}", file=sys.stderr)
        return 1
    rows = parse_folder(folder)
    output = folder / OUTPUT_NAME
    write_workbook(rows, output)
    print(f"Found {len(rows)} payment reference(s) in {folder}")
    print(f"{'reference':<24} {'sent_at':<24} {'response_at':<24} {'status'}")
    for row in rows:
        print(
            f"{row['payment_transaction_reference']:<24} "
            f"{row['sent_at'] or '—':<24} "
            f"{row['response_at'] or '—':<24} "
            f"{row['status']}"
        )
    print(f"Wrote {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
